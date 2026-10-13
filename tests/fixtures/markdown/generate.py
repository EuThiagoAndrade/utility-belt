"""Gera documentos sintéticos próprios, sem dados pessoais ou arquivos do File2MD.

Ferramentas de desenvolvimento: python-docx, python-pptx, openpyxl, xlwt,
Pillow e reportlab. Não são dependências iniciais do Utility Belt.
"""

import json
import struct
import wave
import zipfile
from pathlib import Path

import xlwt
from docx import Document
from openpyxl import Workbook
from PIL import Image, PngImagePlugin
from pptx import Presentation
from reportlab.pdfgen import canvas

ROOT = Path(__file__).parent
TEXT = "Utility Belt fixture"


def make_msg(path):
    """Compound File mínimo com duas propriedades MAPI Unicode em mini-streams."""
    free, end, fat = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD
    header = bytearray(512)
    header[:8] = bytes.fromhex("d0cf11e0a1b11ae1")
    struct.pack_into("<HHHHH", header, 24, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, 1, 1, 0, 4096, 3, 1, end, 0)
    struct.pack_into("<109I", header, 76, 0, *([free] * 108))

    def entry(name, kind, start, size, *, child=free, right=free):
        result = bytearray(128)
        encoded = (name + "\0").encode("utf-16-le")
        result[:len(encoded)] = encoded
        struct.pack_into("<HBBIII", result, 64, len(encoded), kind, 1, free, right, child)
        struct.pack_into("<IQ", result, 116, start, size)
        return result

    subject = TEXT.encode("utf-16-le")
    body = (TEXT + " body").encode("utf-16-le")
    directory = (entry("Root Entry", 5, 2, 128, child=1)
                 + entry("__substg1.0_0037001F", 2, 0, len(subject), right=2)
                 + entry("__substg1.0_1000001F", 2, 1, len(body)) + bytearray(128))
    mini_stream = subject.ljust(64, b"\0") + body.ljust(64, b"\0")
    fat_sector = struct.pack("<128I", fat, end, end, end, *([free] * 124))
    mini_fat = struct.pack("<128I", end, end, *([free] * 126))
    path.write_bytes(header + fat_sector + directory + mini_stream.ljust(512, b"\0") + mini_fat)


def main():
    for name in ("sample.txt", "sample.md", "no-extension", "sample.unrecognized"):
        (ROOT / name).write_text(f"---\ntitle: {TEXT}\n---\n\n{TEXT}\n", encoding="utf-8")
    (ROOT / "sample.json").write_text(json.dumps({"title": TEXT}), encoding="utf-8")
    (ROOT / "sample.jsonl").write_text(json.dumps({"title": TEXT}) + "\n", encoding="utf-8")
    (ROOT / "sample.csv").write_text(f"title,value\n{TEXT},42\n", encoding="utf-8")
    (ROOT / "sample.html").write_text(f"<html><head><title>{TEXT}</title></head><body><h1>{TEXT}</h1></body></html>", encoding="utf-8")
    (ROOT / "sample.rss").write_text(f"<rss version='2.0'><channel><title>{TEXT}</title><link>https://example.com</link><description>Test</description><item><title>{TEXT}</title><description>Test content</description></item></channel></rss>", encoding="utf-8")
    (ROOT / "sample.ipynb").write_text(json.dumps({"nbformat": 4, "nbformat_minor": 5, "metadata": {}, "cells": [{"cell_type": "markdown", "metadata": {}, "source": [TEXT]}]}), encoding="utf-8")
    document = Document()
    document.add_heading(TEXT)
    document.add_paragraph(TEXT)
    document.save(ROOT / "sample.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = TEXT
    presentation.save(ROOT / "sample.pptx")
    workbook = Workbook()
    workbook.active.append([TEXT, 42])
    workbook.save(ROOT / "sample.xlsx")
    old_workbook = xlwt.Workbook()
    old_workbook.add_sheet("Fixture").write(0, 0, TEXT)
    old_workbook.save(str(ROOT / "sample.xls"))
    pdf = canvas.Canvas(str(ROOT / "sample.pdf"), invariant=True)
    pdf.drawString(72, 720, TEXT)
    pdf.save()
    image = Image.new("RGB", (16, 16), color="blue")
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Title", TEXT)
    image.save(ROOT / "sample.png", pnginfo=metadata)
    image.save(ROOT / "sample.jpg")
    with wave.open(str(ROOT / "sample.wav"), "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(16000)
        audio.writeframes(b"\0\0" * 1600)
    make_msg(ROOT / "sample.msg")
    with zipfile.ZipFile(ROOT / "sample.epub", "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
        archive.writestr("content.opf", '<package xmlns="http://www.idpf.org/2007/opf" version="2.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Utility Belt fixture</dc:title></metadata><manifest><item id="c" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c"/></spine></package>')
        archive.writestr("chapter.xhtml", f'<html xmlns="http://www.w3.org/1999/xhtml"><body><h1>{TEXT}</h1></body></html>')
    with zipfile.ZipFile(ROOT / "sample.zip", "w") as archive:
        archive.write(ROOT / "sample.txt", "sub/sample.txt")
        archive.write(ROOT / "sample.pdf", "sub/sample.pdf")
    with zipfile.ZipFile(ROOT / "partial.zip", "w") as archive:
        archive.write(ROOT / "sample.txt", "sample.txt")
        archive.writestr("broken.pdf", b"%PDF-1.4\ncorrupt PDF")


if __name__ == "__main__":
    main()
